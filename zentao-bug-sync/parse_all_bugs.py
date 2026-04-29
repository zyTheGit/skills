import openpyxl
import re

wb = openpyxl.load_workbook(
    r"C:\Users\hpee2\Desktop\zentao-bugs\zentao_bugs_20260415.xlsx", read_only=True
)
ws = wb.active
headers = [cell.value for cell in ws[1]]
bug_id_col = headers.index("BUG ID")
title_col = headers.index("标题")
step_col = headers.index("步骤")
img_col = headers.index("图片链接") if "图片链接" in headers else None

bugs = []
for row in ws.iter_rows(min_row=2):
    bug_id = row[bug_id_col].value
    title = row[title_col].value
    step = row[step_col].value if step_col else ""
    img = row[img_col].value if img_col else ""
    if bug_id:
        bugs.append({"id": str(bug_id), "title": title, "step": step, "img": img})

wb.close()

print(f"总数: {len(bugs)}")
print(f"含图片: {sum(1 for b in bugs if b['img'])}")

with open("bugs_list.txt", "w", encoding="utf-8") as f:
    for b in bugs:
        f.write(f"{b['id']}|{b['title']}|{b['img']}\n")

print("已保存到 bugs_list.txt")
