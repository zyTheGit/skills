import openpyxl

wb = openpyxl.load_workbook(
    r"C:\Users\hpee2\Desktop\zentao-bugs\zentao_bugs_20260415.xlsx"
)
ws = wb.active
headers = [cell.value for cell in ws[1]]
bug_id_col = headers.index("BUG ID")
title_col = headers.index("标题")
step_col = headers.index("步骤")
img_col = headers.index("图片链接")

bugs = []
seen_ids = set()
for row in ws.iter_rows(min_row=2):
    bug_id = str(row[bug_id_col].value)
    if bug_id and bug_id not in seen_ids:
        seen_ids.add(bug_id)
        title = row[title_col].value
        img = row[img_col].value
        bugs.append({"id": bug_id, "title": title, "img": img})

wb.close()

with open("all_bugs.txt", "w", encoding="utf-8") as f:
    for b in bugs:
        f.write(b["id"] + "|" + b["img"] + "\n")

print("共", len(bugs), "个唯一BUG")
