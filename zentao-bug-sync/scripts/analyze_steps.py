#!/usr/bin/env python3
"""本地结构化解析 BUG 步骤，生成根因分析、修复建议、关键操作和环境信息"""

import re
import html
from openpyxl import load_workbook

OPERATION_KEYWORDS = [
    "点击",
    "输入",
    "填写",
    "选择",
    "勾选",
    "上传",
    "下载",
    "删除",
    "修改",
    "保存",
    "提交",
    "刷新",
    "切换",
    "打开",
    "关闭",
    "导出",
    "导入",
    "复制",
    "粘贴",
    "拖拽",
    "滚动",
    "搜索",
    "查询",
    "筛选",
    "排序",
    "登录",
    "登出",
    "注册",
    "审批",
    "审核",
    "启用",
    "禁用",
    "创建",
    "编辑",
    "查看",
    "确认",
    "取消",
    "返回",
    "跳转",
    "点击",
    "click",
    "input",
    "select",
    "upload",
    "delete",
    "save",
    "submit",
    "refresh",
    "login",
    "logout",
]

ENVIRONMENT_KEYWORDS = {
    "浏览器": [
        "浏览器",
        "chrome",
        "firefox",
        "edge",
        "safari",
        "ie",
        "internet explorer",
    ],
    "操作系统": ["操作系统", "windows", "linux", "mac", "ios", "android", "系统版本"],
    "服务器": ["服务器", "server", "主机", "ip", "端口", "地址"],
    "数据库": [
        "数据库",
        "mysql",
        "oracle",
        "sqlserver",
        "postgresql",
        "redis",
        "mongodb",
    ],
    "环境": ["测试环境", "生产环境", "开发环境", "uat", "prod", "dev", "staging"],
    "版本": ["版本", "version", "v", "build", "版本号"],
    "设备": ["设备", "手机", "平板", "电脑", "pc", "mobile", "分辨率"],
    "网络": ["网络", "wifi", "4g", "5g", "vpn", "代理", "防火墙"],
    "账号": ["账号", "用户", "权限", "角色", "admin", "user", "账户"],
}

ROOT_CAUSE_PATTERNS = {
    "数据问题": [
        "数据为空",
        "数据错误",
        "数据重复",
        "数据缺失",
        "数据格式",
        "数据不一致",
        "null",
        "空值",
        "无数据",
        "缺少数据",
    ],
    "接口问题": [
        "接口",
        "api",
        "请求",
        "响应",
        "返回",
        "超时",
        "timeout",
        "404",
        "500",
        "报错",
        "失败",
        "异常",
        "错误码",
    ],
    "界面问题": [
        "显示",
        "展示",
        "样式",
        "布局",
        "排版",
        "对齐",
        "字体",
        "颜色",
        "按钮",
        "弹窗",
        "提示",
        "遮挡",
        "重叠",
    ],
    "功能问题": [
        "功能",
        "无法",
        "不能",
        "不支持",
        "未实现",
        "失效",
        "异常",
        "不生效",
        "不响应",
        "卡死",
        "崩溃",
        "闪退",
    ],
    "逻辑问题": [
        "逻辑",
        "流程",
        "顺序",
        "条件",
        "判断",
        "计算",
        "公式",
        "算法",
        "规则",
        "验证",
        "校验",
    ],
    "性能问题": [
        "慢",
        "卡顿",
        "延迟",
        "加载",
        "响应慢",
        "超时",
        "耗时",
        "内存",
        "cpu",
        "占用",
    ],
    "兼容问题": [
        "兼容",
        "不同浏览器",
        "不同版本",
        "不同设备",
        "不同系统",
        "适配",
    ],
    "权限问题": [
        "权限",
        "无权",
        "禁止",
        "拒绝",
        "无法访问",
        "无权限",
        "权限不足",
    ],
    "配置问题": [
        "配置",
        "设置",
        "参数",
        "选项",
        "默认值",
        "配置项",
    ],
}

FIX_SUGGESTIONS = {
    "数据问题": [
        "检查数据源是否正确",
        "验证数据格式和类型",
        "添加数据完整性校验",
        "检查数据库连接状态",
        "查看数据是否存在空值或异常值",
    ],
    "接口问题": [
        "检查接口请求参数是否正确",
        "验证接口响应数据格式",
        "添加接口异常处理和重试机制",
        "检查接口超时设置",
        "查看服务端日志定位具体错误",
    ],
    "界面问题": [
        "检查 CSS 样式是否正确加载",
        "验证页面布局在不同分辨率下的表现",
        "检查组件渲染逻辑",
        "确认样式兼容性问题",
        "检查前端框架版本兼容性",
    ],
    "功能问题": [
        "检查功能实现代码逻辑",
        "验证功能触发条件是否满足",
        "添加功能异常处理",
        "检查功能依赖的服务是否正常",
        "查看前端控制台错误信息",
    ],
    "逻辑问题": [
        "检查业务逻辑流程是否正确",
        "验证条件判断逻辑",
        "检查计算公式和数据转换",
        "添加逻辑边界条件处理",
        "确认规则配置是否正确",
    ],
    "性能问题": [
        "检查数据库查询效率",
        "添加数据缓存机制",
        "优化接口响应速度",
        "减少不必要的网络请求",
        "检查是否存在死循环或阻塞操作",
    ],
    "兼容问题": [
        "测试不同浏览器兼容性",
        "添加浏览器特性检测",
        "使用跨浏览器兼容的 CSS",
        "检查不同版本的 API 兼容",
        "添加设备适配处理",
    ],
    "权限问题": [
        "检查用户权限配置",
        "验证权限校验逻辑",
        "确认角色分配是否正确",
        "检查接口权限拦截器",
        "查看权限配置表数据",
    ],
    "配置问题": [
        "检查配置文件是否正确",
        "验证配置参数值",
        "确认配置加载逻辑",
        "检查配置项是否存在",
        "添加配置缺失的默认值处理",
    ],
}


def clean_html(html_content):
    """清理 HTML 标签，提取纯文本"""
    if not html_content:
        return ""
    text = re.sub(r"<[^>]+>", "", html_content)
    text = html.unescape(text)
    text = re.sub(r"\s+", " ", text)
    text = text.strip()
    return text


def extract_operations(text):
    """提取关键操作"""
    if not text:
        return []

    operations = []
    sentences = re.split(r"[。\n\r,，；;]", text)

    for sentence in sentences:
        sentence = sentence.strip()
        if len(sentence) < 2:
            continue

        for keyword in OPERATION_KEYWORDS:
            if keyword.lower() in sentence.lower():
                idx = sentence.lower().find(keyword.lower())
                end_idx = min(idx + len(keyword) + 20, len(sentence))
                operation_context = sentence[idx:end_idx].strip()

                if operation_context and operation_context not in operations:
                    operations.append(operation_context)
                break

    return operations[:10]


def extract_environment(text):
    """提取环境信息"""
    if not text:
        return {}

    env_info = {}

    for category, keywords in ENVIRONMENT_KEYWORDS.items():
        found = []
        for keyword in keywords:
            pattern = re.compile(rf"{keyword}[：:\s]*([^\n。,，；]+)", re.IGNORECASE)
            matches = pattern.findall(text)
            for match in matches:
                match = match.strip()
                if match and len(match) < 50 and match not in found:
                    found.append(match)

            if keyword.lower() in text.lower() and not found:
                pattern2 = re.compile(rf"([^\n。,，；]+)\s*{keyword}", re.IGNORECASE)
                matches2 = pattern2.findall(text)
                for match in matches2:
                    match = match.strip()
                    if match and len(match) < 50 and match not in found:
                        found.append(match)

        if found:
            env_info[category] = found[:3]

    return env_info


def analyze_root_cause(text, bug_type=""):
    """分析根因"""
    if not text:
        return []

    causes = []
    text_lower = text.lower()

    for category, patterns in ROOT_CAUSE_PATTERNS.items():
        match_count = 0
        matched_keywords = []
        for pattern in patterns:
            if pattern.lower() in text_lower:
                match_count += 1
                matched_keywords.append(pattern)

        if match_count > 0:
            causes.append(
                {
                    "category": category,
                    "confidence": min(match_count * 30, 90),
                    "keywords": matched_keywords[:5],
                }
            )

    if bug_type:
        bug_type_map = {
            "代码错误": "功能问题",
            "配置": "配置问题",
            "安装部署": "配置问题",
            "安全相关": "权限问题",
            "性能问题": "性能问题",
            "标准规范": "逻辑问题",
            "自动化测试": "功能问题",
        }
        mapped_type = bug_type_map.get(bug_type)
        if mapped_type:
            for cause in causes:
                if cause["category"] == mapped_type:
                    cause["confidence"] += 20

    causes.sort(key=lambda x: x["confidence"], reverse=True)
    return causes[:3]


def generate_fix_suggestions(causes, bug_type=""):
    """生成修复建议"""
    suggestions = []

    for cause in causes:
        category = cause["category"]
        if category in FIX_SUGGESTIONS:
            suggestions.extend(FIX_SUGGESTIONS[category][:3])

    if bug_type and len(suggestions) < 5:
        bug_type_map = {
            "代码错误": "功能问题",
            "配置": "配置问题",
            "安装部署": "配置问题",
            "安全相关": "权限问题",
            "性能问题": "性能问题",
        }
        mapped_type = bug_type_map.get(bug_type)
        if mapped_type and mapped_type in FIX_SUGGESTIONS:
            suggestions.extend(FIX_SUGGESTIONS[mapped_type][:2])

    return suggestions[:8]


def format_analysis_result(operations, env_info, causes, suggestions):
    """格式化分析结果"""
    lines = []

    if operations:
        lines.append("【关键操作】")
        for i, op in enumerate(operations, 1):
            lines.append(f"{i}. {op}")
        lines.append("")

    if env_info:
        lines.append("【环境信息】")
        for category, values in env_info.items():
            values_str = " | ".join(values)
            lines.append(f"- {category}: {values_str}")
        lines.append("")

    if causes:
        lines.append("【根因分析】")
        for cause in causes:
            confidence = cause["confidence"]
            keywords = ", ".join(cause["keywords"][:3])
            lines.append(f"- {cause['category']} (置信度: {confidence}%)")
            if keywords:
                lines.append(f"  匹配关键词: {keywords}")
        lines.append("")

    if suggestions:
        lines.append("【修复建议】")
        for i, sug in enumerate(suggestions, 1):
            lines.append(f"{i}. {sug}")

    return "\n".join(lines)


def analyze_bug_steps(steps_html, bug_type=""):
    """分析 BUG 步骤，返回结构化结果"""
    text = clean_html(steps_html)

    operations = extract_operations(text)
    env_info = extract_environment(text)
    causes = analyze_root_cause(text, bug_type)
    suggestions = generate_fix_suggestions(causes, bug_type)

    result = format_analysis_result(operations, env_info, causes, suggestions)

    return {
        "operations": operations,
        "environment": env_info,
        "root_causes": causes,
        "fix_suggestions": suggestions,
        "formatted_text": result,
    }


def run(input_file, bug_id=None, all_bugs=False):
    """处理 Excel 文件"""
    wb = load_workbook(input_file)
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_map = {name: idx for idx, name in enumerate(headers, 1) if name}

    steps_col = col_map.get("步骤", 14)
    type_col = col_map.get("类型", 6)

    analysis_col = None
    if "步骤解析" in col_map:
        analysis_col = col_map["步骤解析"]
    else:
        analysis_col = ws.max_column + 1
        ws.cell(row=1, column=analysis_col, value="步骤解析")

    if not bug_id and not all_bugs:
        print("请指定 --bug-id <ID> 或 --all")
        return

    analyzed_count = 0

    for row in range(2, ws.max_row + 1):
        bid = ws.cell(row=row, column=col_map.get("BUG ID", 1)).value
        if not bid:
            continue

        if all_bugs or (bug_id and str(bid) == str(bug_id)):
            steps_html = ws.cell(row=row, column=steps_col).value
            bug_type = ws.cell(row=row, column=type_col).value or ""

            result = analyze_bug_steps(steps_html, bug_type)

            ws.cell(row=row, column=analysis_col).value = result["formatted_text"]

            analyzed_count += 1

            if bug_id:
                print(f"\nBUG #{bid} 分析结果:")
                print(result["formatted_text"])
                break

    wb.save(input_file)

    if all_bugs:
        print(f"\n已分析 {analyzed_count} 个 BUG")
        print(f"结果已写入 Excel 的「步骤解析」列")


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="本地结构化解析 BUG 步骤")
    parser.add_argument("--input", type=str, required=True, help="Excel 文件路径")
    parser.add_argument("--bug-id", type=str, help="指定 BUG ID")
    parser.add_argument("--all", action="store_true", help="分析所有 BUG")
    args = parser.parse_args()

    run(args.input, args.bug_id, args.all)
