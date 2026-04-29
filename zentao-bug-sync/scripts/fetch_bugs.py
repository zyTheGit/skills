#!/usr/bin/env python3
"""从禅道拉取 BUG 到 Excel"""

import sys
import argparse
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from zentao_api import ZentaoAPI
import html

SEVERITY_MAP = {1: "1-致命", 2: "2-严重", 3: "3-一般", 4: "4-建议"}
STATUS_MAP = {"active": "激活", "resolved": "已解决", "closed": "已关闭"}
TYPE_MAP = {
    "codeerror": "代码错误",
    "config": "配置",
    "install": "安装部署",
    "security": "安全相关",
    "performance": "性能问题",
    "standard": "标准规范",
    "automation": "自动化测试",
    "trackthings": "跟踪事项",
    "newfeature": "新增需求",
    "designchange": "设计变更",
    "designdefect": "设计缺陷",
    "others": "其他",
}
RESOLUTION_MAP = {
    "bydesign": "设计如此",
    "fixed": "已修复",
    "external": "外部原因",
    "wontfix": "不予解决",
    "postponed": "延期处理",
    "cannotreproduce": "无法重现",
    "tostory": "转为需求",
}
COLUMNS = [
    "BUG ID",
    "标题",
    "状态",
    "严重程度",
    "优先级",
    "类型",
    "指派给",
    "创建人",
    "所属产品",
    "所属模块",
    "激活日期",
    "解决方案",
    "解决版本",
    "步骤",
    "步骤解析",
    "图片链接",
    "期望结果",
    "实际结果",
    "AI 提示语",
    "最后同步时间",
    "同步状态",
]


def extract_images_from_html(html_content, base_url=None):
    """从 HTML 中提取图片链接"""
    if not html_content:
        return []

    pattern = r'<img[^>]+src=["\']([^"\']+)["\']'
    matches = re.findall(pattern, html_content, re.IGNORECASE)

    if base_url:
        full_urls = []
        for img in matches:
            if img.startswith("http"):
                full_url = img
            elif img.startswith("/"):
                full_url = base_url + img
            elif img.startswith("{") and img.endswith("}"):
                file_id = img[1:-1]
                full_url = base_url + "/file-read-" + file_id
            elif not img.startswith("http"):
                full_url = base_url + "/" + img
            else:
                full_url = img
            full_urls.append(full_url)
        return full_urls

    return matches


def extract_text_from_html(html_content):
    """从 HTML 中提取纯文本（去除标签）"""
    if not html_content:
        return ""

    text = re.sub(r"<[^>]+>", "", html_content)
    text = html.unescape(text)
    text = re.sub(r"\s+", " ", text)
    text = text.strip()
    return text


def extract_notes_from_html(html_content):
    """从 HTML 中提取备注信息（查找特定格式的备注）"""
    if not html_content:
        return ""

    notes_patterns = [
        r"\[备注\](.*?)(?:\[|$)",
        r"备注[：:]\s*(.*?)(?:\n|$)",
        r"<p[^>]*>备注[：:]?\s*(.*?)</p>",
    ]

    notes = []
    for pattern in notes_patterns:
        matches = re.findall(pattern, html_content, re.IGNORECASE | re.DOTALL)
        for match in matches:
            note = extract_text_from_html(match).strip()
            if note and note not in notes:
                notes.append(note)

    return "; ".join(notes) if notes else ""


def parse_bug_steps(steps_html, base_url=None):
    """解析 BUG 步骤，提取图片、备注等信息"""
    if not steps_html:
        return "", "", ""

    images = extract_images_from_html(steps_html, base_url)
    notes = extract_notes_from_html(steps_html)
    text = extract_text_from_html(steps_html)

    images_str = "\n".join(images) if images else ""

    return text, notes, images_str


def generate_parse_suggestion(images, notes):
    """生成解析建议"""
    suggestions = []

    if images:
        suggestions.append(f"【发现 {len(images)} 张图片】")
        suggestions.append("解析方式:")
        suggestions.append("1. 在浏览器中打开图片链接直接查看")
        suggestions.append("2. 使用 @image-analyzer skill（如有安装）:")
        for i, img in enumerate(images[:3], 1):
            suggestions.append(
                f'   @image-analyzer {img} "分析BUG截图中的错误信息和界面状态"'
            )
        if len(images) > 3:
            suggestions.append(f"   ... 还有 {len(images) - 3} 张图片")
        suggestions.append("3. 使用 Claude.ai 网页版上传图片进行分析")

        suggestions.append("")
        suggestions.append("图片链接:")
        for img in images[:3]:
            suggestions.append(f"  {img}")

    if notes:
        suggestions.append("")
        suggestions.append(f"【备注信息】")
        suggestions.append(notes[:200])

    return "\n".join(suggestions) if suggestions else ""


def apply_formatting(ws, row_count):
    red_fill = PatternFill("solid", fgColor="FF6B6B")
    orange_fill = PatternFill("solid", fgColor="FFA500")
    blue_font = Font(color="4472C4")
    gray_font = Font(color="808080", italic=True)
    yellow_fill = PatternFill("solid", fgColor="FFFF99")
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    for row in range(2, row_count + 2):
        severity_cell = ws.cell(row=row, column=4)
        status_cell = ws.cell(row=row, column=3)
        sync_cell = ws.cell(row=row, column=len(COLUMNS))

        severity_val = str(severity_cell.value or "")
        status_val = str(status_cell.value or "")

        if "1" in severity_val:
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=row, column=col).fill = red_fill
        elif "2" in severity_val:
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=row, column=col).fill = orange_fill

        if status_val == "已关闭":
            status_cell.font = gray_font
        elif status_val == "已解决":
            status_cell.font = blue_font

        sync_cell.value = "synced"
        sync_cell.font = Font(color="27AE60")

        for col in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col_idx, col_name in enumerate(COLUMNS, 1):
        width = max(15, len(col_name) * 2)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def fetch_bugs(
    product_id=None, status_filter=None, output_file=None, only_my_bugs=True
):
    api = ZentaoAPI()
    api.login()

    print(f"正在连接禅道: {api.base_url}")
    print(f"当前用户: {api.account}")

    users_dict = api.get_users_dict()
    print(f"获取用户映射: {len(users_dict)} 个用户")

    products_dict = api.get_products_dict()
    print(f"获取产品映射: {len(products_dict)} 个产品")

    modules_dict = api.get_modules_dict()
    print(f"获取模块映射: {len(modules_dict)} 个模块")

    if only_my_bugs:
        bugs = api.get_my_bugs(status_filter=status_filter)
        print(f"获取指派给您的 BUG: {len(bugs)} 个")
    else:
        bugs = api.get_bugs(product_id=product_id, status_filter=status_filter)
        print(f"获取到 {len(bugs)} 个 BUG")

    if not bugs:
        print("没有找到 BUG")
        return

    bugs.sort(
        key=lambda b: (
            b.get("severity", 3),
            b.get("id", 0),
        )
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "禅道BUG"

    for col_idx, col_name in enumerate(COLUMNS, 1):
        ws.cell(row=1, column=col_idx, value=col_name)

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    bugs_with_images = 0
    bugs_with_notes = 0
    base_url = api.base_url

    for row_idx, bug in enumerate(bugs, 2):
        ws.cell(row=row_idx, column=1, value=bug.get("id", ""))
        ws.cell(row=row_idx, column=2, value=bug.get("title", ""))
        ws.cell(
            row=row_idx,
            column=3,
            value=STATUS_MAP.get(bug.get("status", ""), bug.get("status", "")),
        )
        ws.cell(
            row=row_idx,
            column=4,
            value=SEVERITY_MAP.get(
                bug.get("severity", 3), str(bug.get("severity", ""))
            ),
        )
        ws.cell(row=row_idx, column=5, value=bug.get("pri", ""))
        ws.cell(
            row=row_idx,
            column=6,
            value=TYPE_MAP.get(bug.get("type", ""), bug.get("type", "")),
        )

        assigned_to = bug.get("assignedTo", bug.get("assigned_to", ""))
        ws.cell(
            row=row_idx,
            column=7,
            value=users_dict.get(assigned_to, assigned_to),
        )

        opened_by = bug.get("openedBy", bug.get("createdBy", ""))
        ws.cell(row=row_idx, column=8, value=users_dict.get(opened_by, opened_by))

        product_id_val = bug.get("product", "")
        product_name = products_dict.get(
            str(product_id_val), bug.get("productName", str(product_id_val))
        )
        ws.cell(row=row_idx, column=9, value=product_name)

        module_id_val = bug.get("module", "")
        module_name = modules_dict.get(
            str(module_id_val), bug.get("moduleName", str(module_id_val))
        )
        ws.cell(row=row_idx, column=10, value=module_name)

        ws.cell(
            row=row_idx,
            column=11,
            value=bug.get("activatedDate", bug.get("createdDate", "")),
        )
        ws.cell(
            row=row_idx,
            column=12,
            value=RESOLUTION_MAP.get(
                bug.get("resolution", ""), bug.get("resolution", "")
            ),
        )
        ws.cell(row=row_idx, column=13, value=bug.get("resolvedBuild", ""))

        steps_html = bug.get("steps", bug.get("stepsToReproduce", ""))
        ws.cell(row=row_idx, column=14, value=steps_html)

        text, notes, images_str = parse_bug_steps(steps_html, base_url)

        images_list = extract_images_from_html(steps_html, base_url)
        suggestion = generate_parse_suggestion(images_list, notes)
        ws.cell(row=row_idx, column=15, value=suggestion)

        ws.cell(row=row_idx, column=16, value=images_str)

        if images_str:
            bugs_with_images += 1
        if notes:
            bugs_with_notes += 1

        ws.cell(row=row_idx, column=17, value=bug.get("expect", ""))
        ws.cell(row=row_idx, column=18, value=bug.get("result", ""))
        ws.cell(row=row_idx, column=19, value="")
        ws.cell(row=row_idx, column=20, value=now_str)
        ws.cell(row=row_idx, column=21, value="synced")

    apply_formatting(ws, len(bugs))

    if not output_file:
        output_file = f"zentao_bugs_{datetime.now().strftime('%Y%m%d')}.xlsx"

    wb.save(output_file)
    print(f"已保存到: {output_file}")
    print(f"共 {len(bugs)} 条记录")
    if bugs_with_images > 0:
        print(f"包含图片的 BUG: {bugs_with_images} 个")
        print(f"图片解析建议已写入 Excel 的「步骤解析」列")
        print()
        print("查看图片方法:")
        print("  1. 打开 Excel，复制「图片链接」列的 URL 在浏览器中查看")
        print("  2. 如安装了 @image-analyzer skill，可使用命令解析图片内容")
        print("  3. 使用 Claude.ai 网页版上传图片进行分析")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="从禅道拉取 BUG 到 Excel")
    parser.add_argument("--product-id", type=int, help="产品 ID")
    parser.add_argument("--output", type=str, help="输出文件名")
    parser.add_argument("--status", type=str, help="状态过滤 (active/resolved/closed)")
    parser.add_argument(
        "--all", action="store_true", help="获取所有 BUG（默认只获取自己的）"
    )
    args = parser.parse_args()
    fetch_bugs(
        product_id=args.product_id,
        status_filter=args.status,
        output_file=args.output,
        only_my_bugs=not args.all,
    )
