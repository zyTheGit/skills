#!/usr/bin/env python3
"""增量同步 Excel 与禅道 BUG"""

import sys
import argparse
from datetime import datetime
from openpyxl import load_workbook
from zentao_api import ZentaoAPI

STATUS_MAP = {"active": "激活", "resolved": "已解决", "closed": "已关闭"}
SEVERITY_MAP = {1: "1-致命", 2: "2-严重", 3: "3-一般", 4: "4-建议"}

COLUMNS = [
    "BUG ID",
    "标题",
    "状态",
    "严重程度",
    "优先级",
    "类型",
    "指派给",
    "创建人",
    "所属产品",
    "所属模块",
    "激活日期",
    "解决方案",
    "解决版本",
    "步骤",
    "期望结果",
    "实际结果",
    "AI 提示语",
    "最后同步时间",
    "同步状态",
]


def sync_bugs(input_file, product_id=None):
    api = ZentaoAPI()
    api.login()

    wb = load_workbook(input_file)
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_map = {name: idx for idx, name in enumerate(headers, 1)}

    existing_ids = set()
    for row in range(2, ws.max_row + 1):
        bid = ws.cell(row=row, column=col_map.get("BUG ID", 1)).value
        if bid:
            existing_ids.add(int(bid))

    bugs = api.get_bugs(product_id=product_id)
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    updated = 0
    new_bugs = 0

    for bug in bugs:
        bid = bug.get("id")
        if not bid:
            continue

        row_data = {
            "BUG ID": bid,
            "标题": bug.get("title", ""),
            "状态": STATUS_MAP.get(bug.get("status", ""), bug.get("status", "")),
            "严重程度": SEVERITY_MAP.get(
                bug.get("severity", 3), str(bug.get("severity", ""))
            ),
            "优先级": bug.get("pri", ""),
            "类型": bug.get("type", ""),
            "指派给": bug.get("assignedTo", bug.get("assigned_to", "")),
            "创建人": bug.get("openedBy", bug.get("createdBy", "")),
            "所属产品": bug.get("product", bug.get("productName", "")),
            "所属模块": bug.get("module", bug.get("moduleName", "")),
            "激活日期": bug.get("activatedDate", bug.get("createdDate", "")),
            "解决方案": bug.get("solution", ""),
            "解决版本": bug.get("resolvedBuild", ""),
            "步骤": bug.get("steps", bug.get("stepsToReproduce", "")),
            "期望结果": bug.get("expect", ""),
            "实际结果": bug.get("result", ""),
            "AI 提示语": "",
            "最后同步时间": now_str,
            "同步状态": "synced",
        }

        if bid in existing_ids:
            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=col_map.get("BUG ID", 1)).value == bid:
                    for col_name, col_idx in col_map.items():
                        if col_name in row_data:
                            ws.cell(row=row, column=col_idx).value = row_data[col_name]
                    ws.cell(
                        row=row, column=col_map.get("最后同步时间", 18)
                    ).value = now_str
                    ws.cell(
                        row=row, column=col_map.get("同步状态", 19)
                    ).value = "synced"
                    updated += 1
                    break
        else:
            next_row = ws.max_row + 1
            for col_idx, col_name in enumerate(COLUMNS, 1):
                ws.cell(row=next_row, column=col_idx, value=row_data.get(col_name, ""))
            new_bugs += 1

    wb.save(input_file)
    print(f"同步完成: 更新 {updated} 条, 新增 {new_bugs} 条")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="增量同步")
    parser.add_argument("--input", type=str, required=True, help="Excel 文件")
    parser.add_argument("--product-id", type=int, help="产品 ID")
    args = parser.parse_args()
    sync_bugs(args.input, args.product_id)
