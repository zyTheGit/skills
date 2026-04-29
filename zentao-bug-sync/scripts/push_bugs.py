#!/usr/bin/env python3
"""将 Excel 中的 BUG 更新推送到禅道"""

import sys
import argparse
from datetime import datetime
from openpyxl import load_workbook
from zentao_api import ZentaoAPI

STATUS_REVERSE = {"激活": "active", "已解决": "resolved", "已关闭": "closed"}
SEVERITY_REVERSE = {"1-致命": 1, "2-严重": 2, "3-一般": 3, "4-建议": 4}


def push_bugs(input_file, dry_run=False):
    api = ZentaoAPI()
    api.login()

    wb = load_workbook(input_file)
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_map = {name: idx for idx, name in enumerate(headers, 1)}

    updated = 0
    skipped = 0
    errors = 0
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for row in range(2, ws.max_row + 1):
        bug_id = ws.cell(row=row, column=col_map.get("BUG ID", 1)).value
        if not bug_id:
            continue

        current_status = ws.cell(row=row, column=col_map.get("同步状态", 19)).value
        if current_status == "synced" and not dry_run:
            skipped += 1
            continue

        updates = {}
        status_val = ws.cell(row=row, column=col_map.get("状态", 3)).value
        if status_val and status_val in STATUS_REVERSE:
            updates["status"] = STATUS_REVERSE[status_val]

        severity_val = ws.cell(row=row, column=col_map.get("严重程度", 4)).value
        if severity_val and severity_val in SEVERITY_REVERSE:
            updates["severity"] = SEVERITY_REVERSE[severity_val]

        pri_val = ws.cell(row=row, column=col_map.get("优先级", 5)).value
        if pri_val:
            updates["pri"] = pri_val

        assigned_val = ws.cell(row=row, column=col_map.get("指派给", 7)).value
        if assigned_val:
            updates["assignedTo"] = assigned_val

        solution_val = ws.cell(row=row, column=col_map.get("解决方案", 12)).value
        if solution_val:
            updates["solution"] = solution_val

        resolved_build_val = ws.cell(row=row, column=col_map.get("解决版本", 13)).value
        if resolved_build_val:
            updates["resolvedBuild"] = resolved_build_val

        if not updates:
            skipped += 1
            continue

        if dry_run:
            print(f"[预览] BUG #{bug_id} 将更新: {updates}")
            updated += 1
            continue

        try:
            if "status" in updates:
                status = updates["status"]
                if status == "resolved":
                    result = api.resolve_bug(
                        bug_id,
                        solution=updates.get("solution", "fixed"),
                        resolved_build=updates.get("resolvedBuild", ""),
                    )
                elif status == "closed":
                    result = api.close_bug(bug_id)
                elif status == "active":
                    result = api.activate_bug(bug_id)
                else:
                    result = api.update_bug(bug_id, updates)
            else:
                result = api.update_bug(bug_id, updates)

            ws.cell(row=row, column=col_map.get("最后同步时间", 18)).value = now_str
            ws.cell(row=row, column=col_map.get("同步状态", 19)).value = "synced"
            updated += 1
            print(f"[成功] BUG #{bug_id} 已更新")
        except Exception as e:
            ws.cell(row=row, column=col_map.get("同步状态", 19)).value = "conflict"
            errors += 1
            print(f"[错误] BUG #{bug_id}: {e}")

    if not dry_run:
        wb.save(input_file)

    print(f"\n完成: 更新 {updated} 条, 跳过 {skipped} 条, 错误 {errors} 条")
    if dry_run:
        print("（预览模式，未实际提交）")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="推送 Excel 更新到禅道")
    parser.add_argument("--input", type=str, required=True, help="Excel 文件")
    parser.add_argument("--dry-run", action="store_true", help="预览模式")
    args = parser.parse_args()
    push_bugs(args.input, args.dry_run)
