#!/usr/bin/env python3
"""为 BUG 生成 AI 分析提示语"""

import sys
import argparse
import re
from openpyxl import load_workbook


def extract_image_links(text):
    """从文本中提取图片链接"""
    if not text:
        return []

    patterns = [
        r'<img[^>]+src=["\']([^"\']+)["\']',
        r'https?://[^\s<>"]+\.(?:png|jpg|jpeg|gif|bmp|webp)',
        r'/[^\s<>"]+\.(?:png|jpg|jpeg|gif|bmp|webp)',
    ]

    links = []
    for pattern in patterns:
        matches = re.findall(pattern, text, re.IGNORECASE)
        links.extend(matches)

    return links


def generate_prompt(bug_data):
    steps = bug_data.get("步骤", "N/A")
    image_links = extract_image_links(steps)

    image_section = ""
    if image_links:
        image_section = f"""
【图片附件】
发现以下图片链接，请使用 @image-analyzer skill 解析图片内容：
{chr(10).join(f"- {link}" for link in image_links)}

对于每张图片，请调用: @image-analyzer {image_links[0]} "描述这个错误截图中的具体问题"
"""

    prompt = f"""你是一个软件测试和质量保障专家。请分析以下 BUG 并提供专业建议。

【BUG 信息】
- BUG ID: {bug_data.get("BUG ID", "N/A")}
- 标题: {bug_data.get("标题", "N/A")}
- 严重程度: {bug_data.get("严重程度", "N/A")}
- 优先级: {bug_data.get("优先级", "N/A")}
- 当前状态: {bug_data.get("状态", "N/A")}
- 类型: {bug_data.get("类型", "N/A")}
- 所属产品: {bug_data.get("所属产品", "N/A")}
- 所属模块: {bug_data.get("所属模块", "N/A")}
- 指派给: {bug_data.get("指派给", "N/A")}
- 创建人: {bug_data.get("创建人", "N/A")}
{image_section}
【重现步骤】
{steps}

【期望结果】
{bug_data.get("期望结果", "N/A")}

【实际结果】
{bug_data.get("实际结果", "N/A")}

【请分析】
1. 根因推测：基于描述推测最可能的根本原因
2. 修复建议：具体的修复方案和技术建议
3. 测试覆盖：需要补充哪些测试场景来防止回归
4. 关联风险：是否可能影响其他模块或功能
5. 优先级评估：当前优先级是否合理，是否需要调整"""
    return prompt


def run(input_file, bug_id=None, all_bugs=False):
    wb = load_workbook(input_file)
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_map = {name: idx for idx, name in enumerate(headers, 1)}

    prompt_col = col_map.get("AI 提示语", 17)

    if not bug_id and not all_bugs:
        print("请指定 --bug-id <ID> 或 --all")
        return

    for row in range(2, ws.max_row + 1):
        bid = ws.cell(row=row, column=col_map.get("BUG ID", 1)).value
        if not bid:
            continue

        if all_bugs or (bug_id and str(bid) == str(bug_id)):
            bug_data = {}
            for name, idx in col_map.items():
                bug_data[name] = ws.cell(row=row, column=idx).value

            prompt = generate_prompt(bug_data)
            ws.cell(row=row, column=prompt_col).value = prompt

            if bug_id:
                print(prompt)
                break

    wb.save(input_file)
    if all_bugs:
        print("已为所有 BUG 生成 AI 提示语")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="生成 AI 提示语")
    parser.add_argument("--input", type=str, required=True, help="Excel 文件")
    parser.add_argument("--bug-id", type=str, help="指定 BUG ID")
    parser.add_argument("--all", action="store_true", help="为所有 BUG 生成")
    args = parser.parse_args()
    run(args.input, args.bug_id, args.all)
